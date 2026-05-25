"""Website access functions for the Day06 antibody assignment."""

from __future__ import annotations

import csv
import gzip
import io
import json
import os
import re
import ssl
import urllib.parse
import urllib.request
from pathlib import Path

from Bio import Entrez
from openpyxl import load_workbook


THERASABDAB_CSV_URL = (
    "https://opig.stats.ox.ac.uk/webapps/sabdab-sabpred/static/downloads/"
    "TheraSAbDab_SeqStruc_OnlineDownload.csv"
)
IEDB_ANTIGEN_SEARCH_URL = "https://query-api.iedb.org/antigen_search"
IEDB_EPITOPE_SEARCH_URL = "https://query-api.iedb.org/epitope_search"
PLABDAB_PAIRED_SEQUENCES_URL = (
    "https://opig.stats.ox.ac.uk/webapps/plabdab/static/downloads/"
    "paired_sequences.csv.gz"
)
LENS_PATSEQ_FILE_ENV = "LENS_PATSEQ_FILE"
PLABDAB_DATA_FILE_ENV = "PLABDAB_DATA_FILE"

Entrez.email = "student@example.com"
Entrez.tool = "python-course-antibody-analysis"
Entrez.max_tries = 1
Entrez.sleep_between_tries = 0

ALIAS_OUTPUT_FILE = Path(__file__).with_name("target_aliases.txt")


def open_url(request: urllib.request.Request, timeout: int = 30):
    """Open a URL, with a certificate fallback for course lab computers."""
    try:
        return urllib.request.urlopen(request, timeout=timeout)
    except urllib.error.URLError as error:
        if "CERTIFICATE_VERIFY_FAILED" not in str(error):
            raise

        insecure_context = ssl._create_unverified_context()
        return urllib.request.urlopen(request, timeout=timeout, context=insecure_context)


def download_therasabdab_table() -> list[dict[str, str]]:
    """Download and parse the public Thera-SAbDab antibody CSV table."""
    request = urllib.request.Request(
        THERASABDAB_CSV_URL,
        headers={"User-Agent": "python-course-antibody-analysis/1.0"},
    )

    with open_url(request, timeout=30) as response:
        csv_text = response.read().decode("utf-8-sig")

    reader = csv.DictReader(io.StringIO(csv_text))
    return list(reader)


def download_plabdab_paired_table() -> list[dict[str, str]]:
    """Download and parse PLAbDab paired antibody sequences."""
    file_path = os.environ.get(PLABDAB_DATA_FILE_ENV)
    if file_path:
        return load_csv_table(Path(file_path))

    request = urllib.request.Request(
        PLABDAB_PAIRED_SEQUENCES_URL,
        headers={"User-Agent": "python-course-antibody-analysis/1.0"},
    )

    with open_url(request, timeout=60) as response:
        with gzip.GzipFile(fileobj=response) as gzip_file:
            csv_text = gzip_file.read().decode("utf-8-sig")

    reader = csv.DictReader(io.StringIO(csv_text))
    return list(reader)


def load_csv_table(file_path: Path) -> list[dict[str, str]]:
    """Load antibody data from a local CSV file."""
    with file_path.open(newline="", encoding="utf-8-sig") as input_file:
        reader = csv.DictReader(input_file)
        return list(reader)


def load_excel_table(file_path: Path) -> list[dict[str, str]]:
    """Load antibody data from a local Excel file."""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(header) if header is not None else "" for header in rows[0]]
        table_rows: list[dict[str, str]] = []
        for row in rows[1:]:
            table_rows.append(
                {
                    headers[index]: "" if value is None else str(value)
                    for index, value in enumerate(row)
                    if index < len(headers) and headers[index]
                }
            )
        return table_rows
    finally:
        workbook.close()


def load_antibody_table(file_path: Path | None = None) -> list[dict[str, str]]:
    """Load antibody rows from a local CSV/XLSX file or download Thera-SAbDab."""
    if file_path is None:
        return download_therasabdab_table()

    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        return load_csv_table(file_path)
    if suffix in {".xlsx", ".xlsm"}:
        return load_excel_table(file_path)
    raise ValueError("Input file must be a CSV or XLSX file.")


def count_pubmed_references(antibody_name: str) -> int:
    """Count PubMed papers mentioning an antibody name using Biopython."""
    search_term = f'"{antibody_name}"[Title/Abstract]'

    insecure_context = ssl._create_unverified_context()
    previous_urlopen = Entrez.urlopen
    Entrez.urlopen = lambda request: urllib.request.urlopen(
        request,
        context=insecure_context,
    )
    try:
        with Entrez.esearch(
            db="pubmed",
            term=search_term,
            retmode="xml",
            retmax=0,
        ) as handle:
            record = Entrez.read(handle)
    finally:
        Entrez.urlopen = previous_urlopen

    return int(record["Count"])


def read_entrez_record(search_function, **kwargs):
    """Run a Biopython Entrez request with SSL handling for course computers."""
    insecure_context = ssl._create_unverified_context()
    previous_urlopen = Entrez.urlopen
    Entrez.urlopen = lambda request: urllib.request.urlopen(
        request,
        context=insecure_context,
    )
    try:
        with search_function(**kwargs) as handle:
            return Entrez.read(handle)
    finally:
        Entrez.urlopen = previous_urlopen


def fetch_ncbi_gene_aliases(
    target_query: str,
    seed_aliases: list[str] | None = None,
    max_genes: int = 5,
) -> list[str]:
    """Fetch human gene aliases from NCBI Gene using Biopython."""
    seed_aliases = seed_aliases or [target_query]
    search_term = f'("{target_query}"[All Fields]) AND "Homo sapiens"[Organism]'
    search_record = read_entrez_record(
        Entrez.esearch,
        db="gene",
        term=search_term,
        retmode="xml",
        retmax=max_genes,
    )

    gene_ids = search_record.get("IdList", [])
    if not gene_ids:
        return seed_aliases

    summary_record = read_entrez_record(
        Entrez.esummary,
        db="gene",
        id=",".join(gene_ids),
        retmode="xml",
    )

    aliases: list[str] = [*seed_aliases]
    for document in summary_record["DocumentSummarySet"]["DocumentSummary"]:
        if not document_matches_seed_aliases(document, seed_aliases):
            continue
        for field in ("Name", "Description", "OtherAliases", "OtherDesignations"):
            value = str(document.get(field, ""))
            if not value:
                continue
            for alias in value.replace("|", ",").split(","):
                alias = alias.strip()
                if alias:
                    aliases.append(alias)

    aliases = list(dict.fromkeys(aliases))
    return aliases or seed_aliases


def fetch_iedb_antigen_aliases(
    target_query: str,
    seed_aliases: list[str] | None = None,
    max_records: int = 10,
) -> list[str]:
    """Fetch antigen names and accessions from IEDB's public query API."""
    seed_aliases = seed_aliases or [target_query]
    aliases: list[str] = [target_query, *seed_aliases]
    records: list[dict] = []

    for seed_alias in seed_aliases:
        records.extend(fetch_iedb_antigen_records(seed_alias, max_records=max_records))

    if not records:
        records.extend(search_iedb_antigen_records(seed_aliases))

    for record in records:
        aliases.extend(extract_iedb_antigen_aliases(record))

    aliases = [alias for alias in aliases if alias.strip()]
    return list(dict.fromkeys(aliases)) or seed_aliases


def fetch_iedb_antigen_records(query: str, max_records: int = 10) -> list[dict]:
    """Fetch exact antigen-name matches from IEDB for one target alias."""
    query = query.strip()
    if not query:
        return []

    params = {
        "limit": str(max_records),
        "select": ",".join(
            [
                "parent_source_antigen_id",
                "parent_source_antigen_names",
                "parent_source_antigen_source_org_name",
                "curated_source_antigens",
            ]
        ),
        "parent_source_antigen_names": f"ov.{{{format_iedb_array_value(query)}}}",
    }
    request = urllib.request.Request(
        f"{IEDB_ANTIGEN_SEARCH_URL}?{urllib.parse.urlencode(params)}",
        headers={"User-Agent": "python-course-antibody-analysis/1.0"},
    )

    with open_url(request, timeout=30) as response:
        data = json.loads(response.read().decode("utf-8"))

    if isinstance(data, list):
        return [record for record in data if isinstance(record, dict)]
    return []


def search_iedb_antigen_records(
    seed_aliases: list[str],
    max_records: int = 5000,
) -> list[dict]:
    """Scan a limited IEDB antigen_search window and keep matching records."""
    params = {
        "limit": str(max_records),
        "select": ",".join(
            [
                "parent_source_antigen_id",
                "parent_source_antigen_names",
                "parent_source_antigen_source_org_name",
                "curated_source_antigens",
            ]
        ),
    }
    request = urllib.request.Request(
        f"{IEDB_ANTIGEN_SEARCH_URL}?{urllib.parse.urlencode(params)}",
        headers={"User-Agent": "python-course-antibody-analysis/1.0"},
    )

    with open_url(request, timeout=60) as response:
        data = json.loads(response.read().decode("utf-8"))

    if not isinstance(data, list):
        return []

    return [
        record
        for record in data
        if isinstance(record, dict) and iedb_record_matches_seed_aliases(record, seed_aliases)
    ]


def search_iedb_epitope_records(
    seed_aliases: list[str],
    max_records: int = 5000,
) -> list[dict]:
    """Scan a limited IEDB epitope_search window and keep matching records."""
    params = {
        "limit": str(max_records),
        "select": ",".join(
            [
                "structure_id",
                "structure_iri",
                "structure_descriptions",
                "structure_type",
                "linear_sequence",
                "reference_ids",
                "curated_source_antigens",
            ]
        ),
    }
    request = urllib.request.Request(
        f"{IEDB_EPITOPE_SEARCH_URL}?{urllib.parse.urlencode(params)}",
        headers={"User-Agent": "python-course-antibody-analysis/1.0"},
    )

    with open_url(request, timeout=60) as response:
        data = json.loads(response.read().decode("utf-8"))

    if not isinstance(data, list):
        return []

    return [
        record
        for record in data
        if isinstance(record, dict) and iedb_record_matches_seed_aliases(record, seed_aliases)
    ]


def format_iedb_array_value(value: str) -> str:
    """Format one string for a PostgREST array overlap filter."""
    if re.fullmatch(r"[A-Za-z0-9_.-]+", value):
        return value

    escaped_value = value.replace("\\", "\\\\").replace('"', '\\"')
    return f'"{escaped_value}"'


def extract_iedb_antigen_aliases(record: dict) -> list[str]:
    """Extract useful target aliases from one IEDB antigen_search record."""
    aliases: list[str] = []
    aliases.extend(list_values(record.get("parent_source_antigen_names")))

    parent_id = str(record.get("parent_source_antigen_id", "")).strip()
    if parent_id:
        aliases.append(parent_id.rsplit("/", 1)[-1])

    curated_antigens = record.get("curated_source_antigens") or []
    if isinstance(curated_antigens, list):
        for antigen in curated_antigens:
            if not isinstance(antigen, dict):
                continue
            aliases.extend(
                [
                    str(antigen.get("name", "")).strip(),
                    str(antigen.get("accession", "")).strip(),
                ]
            )

    return list(dict.fromkeys(alias for alias in aliases if alias))


def iedb_sequence_url(record: dict) -> str:
    """Return a public IEDB epitope page URL when an epitope id is available."""
    structure_id = str(record.get("structure_id", "")).strip()
    if structure_id:
        return f"https://www.iedb.org/epitope/{structure_id}"

    structure_iri = str(record.get("structure_iri", "")).strip()
    match = re.search(r"(\d+)$", structure_iri)
    if match:
        return f"https://www.iedb.org/epitope/{match.group(1)}"
    return ""


def iedb_record_matches_seed_aliases(record: dict, seed_aliases: list[str]) -> bool:
    """Check whether an IEDB record contains one of the target aliases."""
    normalized_seeds = {normalize_alias(alias) for alias in seed_aliases if alias}
    normalized_aliases = {
        normalize_alias(alias)
        for alias in extract_iedb_antigen_aliases(record)
        if alias
    }
    return bool(normalized_seeds & normalized_aliases)


def read_lens_patseq_file() -> list[dict[str, str]]:
    """Read a local Lens PatSeq export file when one is configured."""
    file_path_text = os.environ.get(LENS_PATSEQ_FILE_ENV)
    if not file_path_text:
        return []

    file_path = Path(file_path_text)
    if file_path.suffix.lower() == ".csv":
        return load_csv_table(file_path)
    if file_path.suffix.lower() in {".xlsx", ".xlsm"}:
        return load_excel_table(file_path)
    if file_path.suffix.lower() in {".fasta", ".fa", ".faa"}:
        return read_fasta_records(file_path)
    raise ValueError("Lens PatSeq export must be CSV, XLSX, FASTA, FA, or FAA.")


def read_fasta_records(file_path: Path) -> list[dict[str, str]]:
    """Read simple FASTA records into dictionaries."""
    records: list[dict[str, str]] = []
    current_header = ""
    current_sequence: list[str] = []
    with file_path.open(encoding="utf-8") as input_file:
        for line in input_file:
            line = line.strip()
            if not line:
                continue
            if line.startswith(">"):
                if current_header:
                    records.append(
                        {
                            "id": current_header.split()[0],
                            "description": current_header,
                            "sequence": "".join(current_sequence),
                        }
                    )
                current_header = line[1:]
                current_sequence = []
            else:
                current_sequence.append(line)
    if current_header:
        records.append(
            {
                "id": current_header.split()[0],
                "description": current_header,
                "sequence": "".join(current_sequence),
            }
        )
    return records


def list_values(value) -> list[str]:
    """Return strings from a JSON list or scalar field."""
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    if value:
        return [str(value).strip()]
    return []


def write_target_aliases(target_query: str, aliases: list[str]) -> Path:
    """Save aliases used for the current search to a text file."""
    with ALIAS_OUTPUT_FILE.open("w", encoding="utf-8") as alias_file:
        alias_file.write(f"Target query: {target_query}\n")
        alias_file.write("Aliases used for search:\n")
        for alias in aliases:
            alias_file.write(f"- {alias}\n")
    return ALIAS_OUTPUT_FILE


def normalize_alias(value: str) -> str:
    """Normalize alias text before comparing NCBI Gene records."""
    return re.sub(r"[^a-z0-9]", "", value.lower())


def document_matches_seed_aliases(document, seed_aliases: list[str]) -> bool:
    """Keep only NCBI Gene records that match the user's target or known aliases."""
    normalized_seeds = {normalize_alias(alias) for alias in seed_aliases if alias}
    document_values = []
    for field in ("Name", "Description", "OtherAliases", "OtherDesignations"):
        value = str(document.get(field, ""))
        document_values.extend(value.replace("|", ",").split(","))

    normalized_values = {normalize_alias(value.strip()) for value in document_values if value.strip()}
    return bool(normalized_seeds & normalized_values)
