"""Create a sample-based OCR comparison report and seed personal handwriting data."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.config import Config, PROJECT_ROOT
from src.image_loader import correct_perspective, load_image
from src.learning_store import DEFAULT_LEARNING_FOLDER, load_corrections, save_correction
from src.ocr_engine import OCREngine
from src.review_model import FieldCorrection
from src.table_detection import DEFAULT_CROPS_FOLDER, detect_table
from src.validation import DEFAULT_RANGES


SAMPLE_FOLDER = PROJECT_ROOT.parent / "sample test"
EXPECTED_WORKBOOK = SAMPLE_FOLDER / "Sample.xlsx"
MEASUREMENT_FIELDS = ("W", "L", "Weight")
PAGE_MOUSE_ROWS = 30


@dataclass(frozen=True)
class SampleSet:
    """Map one folder of photographed pages to its verified workbook sheet."""

    name: str
    image_folder: Path
    sheet_name: str


SAMPLE_SETS = (
    SampleSet("pic day 1", SAMPLE_FOLDER / "pic day 1", "output 1"),
    SampleSet("pic day 2", SAMPLE_FOLDER / "pic day 2", "output 2"),
)


@dataclass(frozen=True)
class BaselineSummary:
    """Concise result of a sample baseline run."""

    report_path: Path
    expected_values: int
    exact_matches: int
    personal_matches_after_seeding: int
    saved_learning_examples: int
    skipped_existing_examples: int


def run_sample_baseline() -> BaselineSummary:
    """Compare sample-photo OCR with Sample.xlsx and seed local learning examples."""
    config = Config.load()
    if not EXPECTED_WORKBOOK.is_file():
        raise FileNotFoundError(f"Expected workbook not found: {EXPECTED_WORKBOOK}")
    _validate_sample_sets()

    report_path = _new_output_path(config.output_folder, "sample_baseline_report", ".csv")
    crops_folder = _new_crops_folder(DEFAULT_CROPS_FOLDER)
    engine = OCREngine(enable_personal_handwriting=False)
    existing_keys = _existing_learning_keys()

    rows: list[dict[str, Any]] = []
    exact_matches = 0
    saved_examples = 0
    skipped_existing = 0

    total_expected_values = 0
    for sample_set in SAMPLE_SETS:
        expected_values = _load_expected_values(config, sample_set.sheet_name)
        total_expected_values += sum(
            1
            for values in expected_values.values()
            for field_name in MEASUREMENT_FIELDS
            if values[field_name] is not None
        )
        image_paths = sorted(sample_set.image_folder.glob("*.jpg"))
        for image_index, image_path in enumerate(image_paths):
            image = correct_perspective(load_image(image_path))
            table = detect_table(
                image,
                source_image=image_path.name,
                crops_folder=crops_folder,
                debug=True,
            )
            for page_row in range(1, PAGE_MOUSE_ROWS + 1):
                mouse_id = str(image_index * PAGE_MOUSE_ROWS + page_row)
                if mouse_id not in expected_values:
                    continue
                cells = table.cells_for_row(page_row)
                for field_name in MEASUREMENT_FIELDS:
                    expected = expected_values[mouse_id][field_name]
                    crop_path = cells[field_name].crop_path
                    ocr_result = engine.read_number(crop_path, field_name)
                    matched = _same_number(expected, ocr_result.parsed_value)
                    if expected is not None and matched:
                        exact_matches += 1

                    learning_key = (
                        image_path.name,
                        mouse_id,
                        field_name,
                        _learning_value_key(expected),
                    )
                    saved = False
                    if learning_key in existing_keys:
                        skipped_existing += 1
                    else:
                        saved = save_correction(
                            FieldCorrection(
                                source_image=image_path.name,
                                mouse_id=mouse_id,
                                field_name=field_name,
                                original_prediction=None,
                                corrected_value=expected,
                                crop_path=str(crop_path),
                            )
                        )
                        if saved:
                            saved_examples += 1
                            existing_keys.add(learning_key)

                    rows.append(
                        {
                            "sample_set": sample_set.name,
                            "workbook_sheet": sample_set.sheet_name,
                            "source_image": image_path.name,
                            "page_row": page_row,
                            "mouse_id": mouse_id,
                            "field_name": field_name,
                            "expected_value": _format_value(expected),
                            "ocr_text": ocr_result.text,
                            "ocr_value": _format_value(ocr_result.parsed_value),
                            "ocr_confidence": f"{ocr_result.confidence:.3f}",
                            "exact_match": matched,
                            "learning_example_saved": saved,
                            "crop_path": str(crop_path),
                        }
                    )

    _write_report(report_path, rows)
    personal_matches = _count_personal_matches_after_seeding(rows)
    return BaselineSummary(
        report_path=report_path,
        expected_values=total_expected_values,
        exact_matches=exact_matches,
        personal_matches_after_seeding=personal_matches,
        saved_learning_examples=saved_examples,
        skipped_existing_examples=skipped_existing,
    )


def _load_expected_values(
    config: Config,
    sheet_name: str,
) -> dict[str, dict[str, float | None]]:
    """Read expected sample values from the reference workbook by mouse ID."""
    workbook = load_workbook(EXPECTED_WORKBOOK, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Expected worksheet '{sheet_name}' was not found in {EXPECTED_WORKBOOK}."
            )
        sheet = workbook[sheet_name]
        expected: dict[str, dict[str, float | None]] = {}
        for row in range(config.mouse_id_start_row, sheet.max_row + 1):
            raw_mouse_id = sheet.cell(row, 2).value
            if raw_mouse_id is None:
                continue
            mouse_id = str(int(raw_mouse_id)) if isinstance(raw_mouse_id, float) else str(raw_mouse_id)
            expected[mouse_id] = {
                "W": _optional_number(sheet.cell(row, 11).value, "W"),
                "L": _optional_number(sheet.cell(row, 12).value, "L"),
                "Weight": _optional_number(
                    sheet.cell(row, 13).value,
                    "Weight",
                ),
            }
        return expected
    finally:
        workbook.close()


def _validate_sample_sets() -> None:
    """Validate the expected two-day folder and workbook structure before writing."""
    workbook = load_workbook(EXPECTED_WORKBOOK, read_only=True)
    try:
        for sample_set in SAMPLE_SETS:
            if not sample_set.image_folder.is_dir():
                raise FileNotFoundError(
                    f"Sample image folder not found: {sample_set.image_folder}"
                )
            image_paths = sorted(sample_set.image_folder.glob("*.jpg"))
            if len(image_paths) != 3:
                raise FileNotFoundError(
                    f"Expected exactly 3 JPG images in {sample_set.image_folder}; "
                    f"found {len(image_paths)}."
                )
            if sample_set.sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Expected worksheet '{sample_set.sheet_name}' was not found in "
                    f"{EXPECTED_WORKBOOK}."
                )
    finally:
        workbook.close()


def _existing_learning_keys() -> set[tuple[str, str, str, str]]:
    """Return only the latest saved label for each source mouse and field."""
    latest_corrections = {}
    for correction in load_corrections(DEFAULT_LEARNING_FOLDER):
        if correction.crop_path is None:
            continue
        if not Path(correction.crop_path).is_file():
            continue
        latest_corrections[
            (
                correction.source_image,
                correction.mouse_id,
                correction.field_name,
            )
        ] = correction
    return {
        (
            correction.source_image,
            correction.mouse_id,
            correction.field_name,
            _learning_value_key(correction.corrected_value),
        )
        for correction in latest_corrections.values()
    }


def _learning_value_key(value: float | None) -> str:
    """Create a stable duplicate key for numeric and explicitly empty labels."""
    return "<EMPTY>" if value is None else f"{float(value):.2f}"


def _new_crops_folder(root: Path) -> Path:
    """Create a unique crop folder for one baseline run."""
    for attempt in range(100):
        extra = f"_{attempt:02d}" if attempt else ""
        candidate = root / f"sample_baseline_{datetime.now():%Y%m%d_%H%M%S_%f}{extra}"
        if not candidate.exists():
            candidate.mkdir(parents=True)
            return candidate
    raise FileExistsError(f"Could not create a unique crop folder under {root}")


def _new_output_path(folder: Path, stem: str, suffix: str) -> Path:
    """Create a timestamped path without overwriting an existing report."""
    for attempt in range(100):
        extra = f"_{attempt:02d}" if attempt else ""
        candidate = folder / f"{stem}_{datetime.now():%Y%m%d_%H%M%S_%f}{extra}{suffix}"
        if not candidate.exists():
            return candidate
    raise FileExistsError(f"Could not create a unique output file in {folder}")


def _write_report(path: Path, rows: list[dict[str, Any]]) -> None:
    """Write the baseline comparison CSV."""
    if not rows:
        raise ValueError("No comparison rows were generated.")
    with path.open("w", newline="", encoding="utf-8") as report_file:
        writer = csv.DictWriter(report_file, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _count_personal_matches_after_seeding(rows: list[dict[str, Any]]) -> int:
    """Count matches after the sample examples are available to the recognizer."""
    engine = OCREngine()
    matches = 0
    for row in rows:
        if not row["expected_value"]:
            continue
        result = engine.read_number(row["crop_path"], row["field_name"])
        if _same_number(float(row["expected_value"]), result.parsed_value):
            matches += 1
    return matches


def _optional_number(value: Any, field_name: str) -> float | None:
    """Validate workbook labels before they can enter the learning dataset."""
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Expected a numeric or empty workbook cell, got {value!r}")
    number = float(value)
    minimum, maximum = DEFAULT_RANGES[field_name]
    if not minimum <= number <= maximum:
        raise ValueError(
            f"Invalid {field_name} learning value {number:.2f}; expected "
            f"{minimum:.2f}-{maximum:.2f}. Correct the source workbook first."
        )
    return number


def _same_number(expected: float | None, actual: float | None) -> bool:
    """Compare optional numbers after rounding to workbook precision."""
    if expected is None or actual is None:
        return expected is actual
    return round(float(expected), 2) == round(float(actual), 2)


def _format_value(value: float | None) -> str:
    """Format optional numbers for the CSV report."""
    return "" if value is None else f"{float(value):.2f}"


def main() -> None:
    """Run the sample baseline and print a short summary."""
    summary = run_sample_baseline()
    print(f"Report: {summary.report_path}")
    print(f"Expected non-empty values: {summary.expected_values}")
    print(f"Exact OCR matches before sample seeding: {summary.exact_matches}")
    print(
        "Personal handwriting matches after sample seeding: "
        f"{summary.personal_matches_after_seeding}"
    )
    print(f"Saved learning examples: {summary.saved_learning_examples}")
    print(f"Skipped existing learning examples: {summary.skipped_existing_examples}")


if __name__ == "__main__":
    main()
