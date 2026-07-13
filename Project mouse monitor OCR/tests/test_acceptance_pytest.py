"""Pytest acceptance coverage for configuration, validation, mapping, and export."""

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
import pytest

from src.config import Config, DEFAULT_SETTINGS_PATH
from src.excel_template import DuplicateMouseIDError, build_mouse_id_mapping
from src.export_excel import (
    DuplicateMeasurementIDError,
    export_validated_measurements_to_excel,
)
from src.review_model import ExtractedMeasurement
from src.validation import ValidationStatus, validate_measurement


pytestmark = pytest.mark.full


@pytest.fixture()
def configured_project(tmp_path: Path) -> tuple[Config, Path]:
    """Create a realistic isolated project layout and structure-rich workbook."""
    config_folder = tmp_path / "config"
    templates_folder = tmp_path / "data" / "templates"
    output_folder = tmp_path / "data" / "output"
    images_folder = tmp_path / "data" / "input_images"
    config_folder.mkdir()
    templates_folder.mkdir(parents=True)
    output_folder.mkdir()
    images_folder.mkdir()

    template_path = templates_folder / "Book1.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "Cage A"
    sheet["B2"] = "001"
    sheet["D2"] = "=1+1"
    sheet["E2"].fill = PatternFill(fill_type="solid", fgColor="00FF00")
    sheet["B5"] = 2
    sheet["A7"] = "Cage B"
    sheet.row_dimensions[4].height = 24
    sheet.merge_cells("A7:C7")
    workbook.save(template_path)
    workbook.close()

    settings = json.loads(DEFAULT_SETTINGS_PATH.read_text(encoding="utf-8"))
    settings["sheet_name"] = "Sheet1"
    settings["mouse_id_start_row"] = 1
    settings_path = config_folder / "settings.json"
    settings_path.write_text(json.dumps(settings), encoding="utf-8")
    return Config.load(settings_path), template_path


def test_load_config(configured_project: tuple[Config, Path]) -> None:
    """1. Load central settings and resolve their required paths."""
    config, template = configured_project

    assert config.excel_template_path == template.resolve()
    assert config.mouse_id_column == "B"
    assert config.output_columns == {"W": "K", "L": "L", "Weight": "M"}


def test_map_mouse_ids_from_template(
    configured_project: tuple[Config, Path],
) -> None:
    """2. Map mouse IDs to original, non-compressed Excel row numbers."""
    config, template = configured_project
    workbook = load_workbook(template)

    metadata = build_mouse_id_mapping(
        workbook, config.sheet_name, config.mouse_id_column
    )

    assert metadata.mouse_id_to_row == {"001": 2, "2": 5}
    assert metadata.mapped_mouse_count == 2
    workbook.close()


def test_duplicate_mouse_ids_in_template_raise_clear_error() -> None:
    """3. Reject a template containing duplicate normalized IDs."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["B2"] = 12
    sheet["B9"] = 12.0

    with pytest.raises(DuplicateMouseIDError, match="rows 2 and 9"):
        build_mouse_id_mapping(workbook, "Sheet1", "B")


def test_zero_is_valid_for_every_measurement() -> None:
    """4. Preserve 0.00 as valid for W, L, and Weight."""
    row = ExtractedMeasurement(source_image="page.jpg", W=0.0, L=0.0, Weight=0.0)

    result = validate_measurement(row)

    assert result.status is ValidationStatus.OK
    assert result.export_allowed
    assert (row.W, row.L, row.Weight) == (0.0, 0.0, 0.0)


@pytest.mark.parametrize("missing_field", ["W", "L"])
def test_missing_w_or_l_requires_approval(missing_field: str) -> None:
    """5. Block a row with missing W/L until explicit row approval."""
    values = {"W": 1.0, "L": 2.0, "Weight": 20.0}
    values[missing_field] = None
    row = ExtractedMeasurement(source_image="page.jpg", **values)

    result = validate_measurement(row)

    assert result.status is ValidationStatus.REQUIRES_APPROVAL
    assert not result.export_allowed
    row.approved = True
    assert result.export_allowed


def test_missing_weight_warns_without_blocking() -> None:
    """6. Highlight missing Weight as a non-blocking warning."""
    row = ExtractedMeasurement(source_image="page.jpg", W=1.0, L=2.0, Weight=None)

    result = validate_measurement(row)

    assert result.status is ValidationStatus.WARNING
    assert result.export_allowed
    assert "Weight is missing." in result.warnings


@pytest.mark.parametrize(
    ("field_name", "value"),
    [("W", 20.01), ("L", -0.01), ("Weight", 40.01)],
)
def test_out_of_range_values_require_approval(field_name: str, value: float) -> None:
    """7. Require approval for negative or excessive measurements."""
    values = {"W": 1.0, "L": 2.0, "Weight": 20.0}
    values[field_name] = value
    row = ExtractedMeasurement(source_image="page.jpg", **values)

    result = validate_measurement(row)

    assert result.status is ValidationStatus.REQUIRES_APPROVAL
    assert result.requires_explicit_approval
    assert not result.export_allowed


def test_export_writes_only_klm_and_preserves_other_cells_and_empty_rows(
    configured_project: tuple[Config, Path],
) -> None:
    """8-10. Write K/L/M only and preserve formulas, styles, merges, and blank rows."""
    config, template = configured_project
    output = config.output_folder / "acceptance.xlsx"
    measurements = [
        ExtractedMeasurement(
            source_image="page.jpg",
            mouse_id="001",
            W=0.0,
            L=2.5,
            Weight=21.25,
            approved=True,
        ),
        ExtractedMeasurement(
            source_image="page.jpg",
            mouse_id="2",
            W=9.0,
            L=9.0,
            Weight=9.0,
            approved=False,
        ),
    ]

    export_validated_measurements_to_excel(template, output, measurements, config)

    original = load_workbook(template, data_only=False)
    exported = load_workbook(output, data_only=False)
    source_sheet = original["Sheet1"]
    output_sheet = exported["Sheet1"]
    assert output_sheet["K2"].value == 0.0
    assert output_sheet["L2"].value == 2.5
    assert output_sheet["M2"].value == 21.25
    assert output_sheet["K2"].number_format == "0.00"
    assert output_sheet["K5"].value is None
    assert output_sheet["D2"].value == "=1+1"
    assert output_sheet["E2"].fill.fgColor.rgb == source_sheet["E2"].fill.fgColor.rgb
    assert output_sheet.row_dimensions[4].height == 24
    assert {str(item) for item in output_sheet.merged_cells.ranges} == {
        str(item) for item in source_sheet.merged_cells.ranges
    }
    assert all(output_sheet.cell(4, column).value is None for column in range(1, 14))

    allowed_changes = {"K2", "L2", "M2"}
    coordinates = set(source_sheet._cells) | set(output_sheet._cells)
    for row_number, column_number in coordinates:
        coordinate = output_sheet.cell(row_number, column_number).coordinate
        if coordinate not in allowed_changes:
            assert output_sheet.cell(row_number, column_number).value == source_sheet.cell(
                row_number, column_number
            ).value
    original.close()
    exported.close()


def test_duplicate_extracted_mouse_ids_block_export_by_default(
    configured_project: tuple[Config, Path],
) -> None:
    """11. Stop before creating output when extracted measurements duplicate an ID."""
    config, template = configured_project
    output = config.output_folder / "must_not_exist.xlsx"
    measurements = [
        ExtractedMeasurement(source_image="one.jpg", mouse_id="001", approved=True),
        ExtractedMeasurement(source_image="two.jpg", mouse_id="001", approved=True),
    ]

    with pytest.raises(DuplicateMeasurementIDError, match="Export was stopped"):
        export_validated_measurements_to_excel(template, output, measurements, config)
    assert not output.exists()
