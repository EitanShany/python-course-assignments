"""End-to-end integration checks against the user-provided sample artifacts."""

from dataclasses import replace
from pathlib import Path
import shutil

import cv2
from openpyxl import load_workbook
import pytest

from src.config import Config, PROJECT_ROOT
from src.excel_template import inspect_template
from src.export_excel import export_validated_measurements_to_excel
from src.ocr_engine import OCREngine
from src.review_model import ExtractedMeasurement
from src.table_detection import detect_table


pytestmark = pytest.mark.full


SAMPLE_FOLDER = PROJECT_ROOT.parent / "sample test"
EXPECTED_WORKBOOK = SAMPLE_FOLDER / "Sample.xlsx"
SAMPLE_SETS = (
    (SAMPLE_FOLDER / "pic day 1", "output 1"),
    (SAMPLE_FOLDER / "pic day 2", "output 2"),
)
SAMPLE_IMAGES = [
    image_path
    for image_folder, _ in SAMPLE_SETS
    for image_path in sorted(image_folder.glob("*.jpg"))
]


def _require_samples() -> None:
    """Skip cleanly when the optional user sample folder is not distributed."""
    if not EXPECTED_WORKBOOK.is_file() or any(
        len(list(image_folder.glob("*.jpg"))) != 3
        for image_folder, _ in SAMPLE_SETS
    ):
        pytest.skip("User-provided sample images/workbook are not available.")
    workbook = load_workbook(EXPECTED_WORKBOOK, read_only=True)
    try:
        if any(sheet_name not in workbook.sheetnames for _, sheet_name in SAMPLE_SETS):
            pytest.skip("Expected sample workbook sheets are not available.")
    finally:
        workbook.close()


def test_default_template_matches_sample_mouse_layout() -> None:
    """Use the real output sheet and map exactly mouse IDs 1 through 70."""
    _require_samples()
    config = Config.load()

    metadata = inspect_template(config)

    assert metadata.sheet_name == "output"
    assert metadata.mapped_mouse_count == 70
    assert metadata.mouse_id_to_row["1"] == 9
    assert metadata.mouse_id_to_row["70"] == 78
    workbook = load_workbook(config.excel_template_path)
    sheet = workbook[config.sheet_name]
    assert all(
        sheet.cell(row, column).value is None
        for row in range(9, 79)
        for column in range(11, 14)
    )
    workbook.close()


@pytest.mark.parametrize("expected_sheet_name", ["output 1", "output 2"])
def test_export_reproduces_sample_expected_values(
    tmp_path: Path,
    expected_sheet_name: str,
) -> None:
    """Recreate each day's expected K/L/M values in the unchanged output template."""
    _require_samples()
    config = replace(Config.load(), output_folder=tmp_path)
    expected = load_workbook(EXPECTED_WORKBOOK, data_only=False)
    expected_sheet = expected[expected_sheet_name]
    measurements: list[ExtractedMeasurement] = []
    for row in range(config.mouse_id_start_row, expected_sheet.max_row + 1):
        mouse_id = expected_sheet.cell(row, 2).value
        values = [expected_sheet.cell(row, column).value for column in range(11, 14)]
        if mouse_id is None or not any(value is not None for value in values):
            continue
        measurements.append(
            ExtractedMeasurement(
                source_image="sample_reference",
                mouse_id=str(mouse_id),
                W=values[0],
                L=values[1],
                Weight=values[2],
                approved=True,
            )
        )

    output_path = tmp_path / f"{expected_sheet_name.replace(' ', '_')}.xlsx"
    result = export_validated_measurements_to_excel(
        config.excel_template_path,
        output_path,
        measurements,
        config,
    )

    actual = load_workbook(output_path, data_only=False)
    actual_sheet = actual[config.sheet_name]
    assert result.written_measurements == len(measurements)
    assert result.warnings == []
    for row in range(config.mouse_id_start_row, expected_sheet.max_row + 1):
        for column in range(11, 14):
            assert actual_sheet.cell(row, column).value == expected_sheet.cell(
                row, column
            ).value
    actual.close()
    expected.close()


def test_sample_photos_produce_aligned_cell_sets(tmp_path: Path) -> None:
    """Detect all six photographed pages, exclude footers, and save 150 cells each."""
    _require_samples()
    for image_path in SAMPLE_IMAGES:
        image = cv2.imread(str(image_path))
        assert image is not None

        result = detect_table(
            image,
            source_image=image_path.name,
            crops_folder=tmp_path,
            debug=True,
        )

        assert result.table_detection_method == "detected_grid"
        assert len(result.cells) == 150
        assert result.cells[-1].bbox.y2 <= result.table_bbox.y2
        assert result.debug_image_path is not None
        assert result.debug_image_path.is_file()
        assert all(cell.crop_path and cell.crop_path.is_file() for cell in result.cells)


def test_sample_printed_mouse_id_ocr_benchmark(tmp_path: Path) -> None:
    """Track printed mouse-ID OCR accuracy on the sample page crops."""
    _require_samples()
    pytest.importorskip("pytesseract")
    if shutil.which("tesseract") is None:
        pytest.skip("Tesseract executable is not available on PATH.")

    engine = OCREngine(enable_personal_handwriting=False)
    expected_total = 0
    exact_matches = 0
    readable = 0
    mismatches: list[str] = []
    for image_folder, _ in SAMPLE_SETS:
        for page_index, image_path in enumerate(sorted(image_folder.glob("*.jpg"))):
            image = cv2.imread(str(image_path))
            assert image is not None
            result = detect_table(
                image,
                source_image=image_path.name,
                crops_folder=tmp_path,
            )
            for page_row in range(1, 31):
                expected_mouse_id = page_index * 30 + page_row
                if expected_mouse_id > 70:
                    continue
                cells = result.cells_for_row(page_row)
                ocr_result = engine.read_printed_mouse_id(cells["mouse_id"].crop_path)
                expected_total += 1
                if ocr_result.text:
                    readable += 1
                if ocr_result.text == str(expected_mouse_id):
                    exact_matches += 1
                elif len(mismatches) < 8:
                    mismatches.append(
                        f"{image_path.name} row {page_row}: expected {expected_mouse_id}, got {ocr_result.text or '(blank)'}"
                    )

    assert expected_total == 140
    assert readable / expected_total >= 0.35, mismatches
    assert exact_matches / expected_total >= 0.25, mismatches
