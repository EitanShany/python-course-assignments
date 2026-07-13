"""Tests for structure-preserving Excel export."""

from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from src.config import Config
from src.export_excel import (
    DuplicateMeasurementIDError,
    export_validated_measurements_to_excel,
)
from src.review_model import ExtractedMeasurement


class ExportExcelTests(unittest.TestCase):
    """Verify approved-only writes and workbook structure preservation."""

    def setUp(self) -> None:
        """Create an isolated template and matching configuration."""
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        templates = self.root / "data" / "templates"
        output = self.root / "data" / "output"
        images = self.root / "data" / "input_images"
        templates.mkdir(parents=True)
        output.mkdir()
        images.mkdir()
        self.template = templates / "Book1.xlsx"

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet["A1"] = "Cage A"
        sheet["B2"] = "001"
        sheet["D2"] = "=1+1"
        sheet["M2"] = "existing"
        sheet["B4"] = 2
        sheet["K4"] = "unchanged"
        sheet["B7"] = 3
        sheet["E5"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        sheet.merge_cells("A6:C6")
        sheet["A6"] = "Cage B"
        sheet.row_dimensions[3].height = 25
        workbook.save(self.template)
        workbook.close()

        self.config = Config(
            project_root=self.root,
            excel_template_path=self.template,
            output_folder=output,
            input_images_folder=images,
            sheet_name="Sheet1",
            mouse_id_column="B",
            output_columns={"W": "K", "L": "L", "Weight": "M"},
            valid_ranges={"W": (0, 20), "L": (0, 20), "Weight": (0, 40)},
            excel_rules={},
            mapping_key="mouse_id",
            zero_is_valid=True,
            decimal_separator=".",
            decimal_places=2,
            excel_number_format="0.00",
        )

    def test_writes_only_approved_matching_values(self) -> None:
        """Write K/L/M by mouse ID while retaining blanks, formulas, and formatting."""
        records = [
            ExtractedMeasurement(
                source_image="page.jpg", mouse_id="001", W=0, L=2.5, Weight=None, approved=True
            ),
            ExtractedMeasurement(
                source_image="page.jpg", mouse_id="2", W=9, L=9, Weight=9, approved=False
            ),
            ExtractedMeasurement(
                source_image="page.jpg", mouse_id="missing", W=1, L=1, Weight=1, approved=True
            ),
        ]
        destination = self.config.output_folder / "result.xlsx"

        result = export_validated_measurements_to_excel(
            self.template, destination, records, self.config
        )

        exported = load_workbook(destination, data_only=False)
        sheet = exported["Sheet1"]
        self.assertEqual(sheet["K2"].value, 0.0)
        self.assertEqual(sheet["L2"].value, 2.5)
        self.assertEqual(sheet["M2"].value, "existing")
        self.assertEqual(sheet["K2"].number_format, "0.00")
        self.assertEqual(sheet["L2"].number_format, "0.00")
        self.assertEqual(sheet["K4"].value, "unchanged")
        self.assertEqual(sheet["D2"].value, "=1+1")
        self.assertEqual(sheet.row_dimensions[3].height, 25)
        self.assertIn("A6:C6", {str(item) for item in sheet.merged_cells.ranges})
        self.assertEqual(sheet["E5"].fill.fgColor.rgb, "00FFFF00")
        exported.close()

        original = load_workbook(self.template, data_only=False)
        self.assertIsNone(original["Sheet1"]["K2"].value)
        original.close()
        self.assertEqual(result.written_measurements, 1)
        self.assertEqual(len(result.warnings), 1)
        self.assertIn("not found", result.warnings[0])

    def test_duplicate_measurements_stop_before_file_is_created(self) -> None:
        """Reject duplicate IDs without producing a partial workbook."""
        records = [
            ExtractedMeasurement(source_image="a.jpg", mouse_id="2", approved=True),
            ExtractedMeasurement(source_image="b.jpg", mouse_id="2.0", approved=True),
        ]
        destination = self.config.output_folder / "duplicate.xlsx"

        # Numeric IDs normalize consistently when supplied as numeric values.
        records[1].mouse_id = 2.0  # type: ignore[assignment]
        with self.assertRaises(DuplicateMeasurementIDError):
            export_validated_measurements_to_excel(
                self.template, destination, records, self.config
            )
        self.assertFalse(destination.exists())

    def test_refuses_to_overwrite_template(self) -> None:
        """Never save export data over the source workbook."""
        with self.assertRaisesRegex(ValueError, "must not overwrite"):
            export_validated_measurements_to_excel(
                self.template, self.template, [], self.config
            )

    def test_duplicate_override_uses_first_approved_occurrence(self) -> None:
        """Allow the reviewer to approve a later duplicate when the first is rejected."""
        records = [
            ExtractedMeasurement(
                source_image="first.jpg", mouse_id="001", W=9, approved=False
            ),
            ExtractedMeasurement(
                source_image="second.jpg", mouse_id="001", W=3.25, approved=True
            ),
        ]
        destination = self.config.output_folder / "override.xlsx"

        result = export_validated_measurements_to_excel(
            self.template,
            destination,
            records,
            self.config,
            allow_duplicate_measurement_ids=True,
        )

        exported = load_workbook(destination)
        self.assertEqual(exported["Sheet1"]["K2"].value, 3.25)
        exported.close()
        self.assertEqual(result.written_measurements, 1)

    def test_approved_out_of_range_value_is_written_and_absent_mouse_stays_empty(self) -> None:
        """Honor explicit approval and leave template mice with no extracted data untouched."""
        destination = self.config.output_folder / "approved_flag.xlsx"
        record = ExtractedMeasurement(
            source_image="page.jpg",
            mouse_id="001",
            W=25.0,
            L=2.0,
            Weight=45.0,
            approved=True,
        )

        export_validated_measurements_to_excel(
            self.template, destination, [record], self.config
        )

        exported = load_workbook(destination)
        sheet = exported["Sheet1"]
        self.assertEqual(sheet["K2"].value, 25.0)
        self.assertEqual(sheet["M2"].value, 45.0)
        self.assertIsNone(sheet["K7"].value)
        self.assertIsNone(sheet["L7"].value)
        self.assertIsNone(sheet["M7"].value)
        exported.close()


if __name__ == "__main__":
    unittest.main()
