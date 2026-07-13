"""Read the Excel template and map mouse IDs without modifying the workbook."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.workbook.workbook import Workbook

from .config import Config


MAX_XLSX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024


class DuplicateMouseIDError(ValueError):
    """Raised when one mouse ID occurs on more than one template row."""


@dataclass(frozen=True)
class TemplateMetadata:
    """Describe the populated worksheet and its mouse-to-row mapping."""

    sheet_name: str
    mapped_mouse_count: int
    first_used_row: int | None
    last_used_row: int | None
    mouse_id_to_row: dict[str, int]


def validate_xlsx_upload(content: bytes) -> None:
    """Validate size and required ZIP members before persisting an uploaded workbook."""
    if not content:
        raise ValueError("Uploaded Excel template is empty.")
    if len(content) > MAX_XLSX_UPLOAD_BYTES:
        raise ValueError(
            f"Uploaded Excel template exceeds the "
            f"{MAX_XLSX_UPLOAD_BYTES // (1024 * 1024)} MB limit."
        )
    try:
        with ZipFile(BytesIO(content)) as archive:
            names = set(archive.namelist())
            required = {"[Content_Types].xml", "xl/workbook.xml"}
            if not required.issubset(names):
                raise ValueError("Uploaded file is not a valid XLSX workbook.")
            uncompressed_size = sum(item.file_size for item in archive.infolist())
            if uncompressed_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValueError("Uploaded XLSX expands beyond the 200 MB safety limit.")
    except BadZipFile as exc:
        raise ValueError("Uploaded file is not a valid XLSX ZIP archive.") from exc


def load_template(path: Path | str) -> Workbook:
    """Open an existing workbook with formulas retained and editing disabled by convention."""
    template_path = Path(path)
    if not template_path.is_file():
        raise FileNotFoundError(f"Excel template not found: {template_path.resolve()}")
    return load_workbook(filename=template_path, data_only=False, read_only=False)


def build_mouse_id_mapping(
    workbook: Workbook,
    sheet_name: str,
    mouse_id_column: str,
    first_data_row: int = 1,
) -> TemplateMetadata:
    """Read mouse IDs and worksheet bounds without changing any workbook cell."""
    if sheet_name not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames) or "none"
        raise ValueError(
            f"Worksheet '{sheet_name}' was not found. Available sheets: {available}."
        )

    try:
        column_number = column_index_from_string(mouse_id_column.strip().upper())
    except (AttributeError, ValueError) as exc:
        raise ValueError(
            f"Invalid mouse ID column '{mouse_id_column}'. Expected an Excel column such as B."
        ) from exc
    if not isinstance(first_data_row, int) or isinstance(first_data_row, bool) or first_data_row < 1:
        raise ValueError("first_data_row must be a positive integer.")

    worksheet = workbook[sheet_name]
    mouse_id_to_row: dict[str, int] = {}

    # Read only cells already present in the file. worksheet.cell()/iter_rows()
    # can materialize blank cells and would therefore mutate the in-memory sheet.
    id_cells = sorted(
        (
            cell
            for cell in worksheet._cells.values()
            if cell.column == column_number and cell.row >= first_data_row
        ),
        key=lambda cell: cell.row,
    )
    for cell in id_cells:
        row_number = cell.row
        mouse_id = normalize_mouse_id(cell.value)
        if mouse_id is None:
            continue
        previous_row = mouse_id_to_row.get(mouse_id)
        if previous_row is not None:
            raise DuplicateMouseIDError(
                f"Duplicate mouse ID '{mouse_id}' in worksheet '{sheet_name}': "
                f"rows {previous_row} and {row_number}."
            )
        mouse_id_to_row[mouse_id] = row_number

    first_used_row, last_used_row = _find_used_row_bounds(worksheet)
    return TemplateMetadata(
        sheet_name=worksheet.title,
        mapped_mouse_count=len(mouse_id_to_row),
        first_used_row=first_used_row,
        last_used_row=last_used_row,
        mouse_id_to_row=mouse_id_to_row,
    )


def inspect_template(config: Config) -> TemplateMetadata:
    """Open the configured template and return its read-only inspection metadata."""
    workbook = load_template(config.excel_template_path)
    try:
        return build_mouse_id_mapping(
            workbook,
            sheet_name=config.sheet_name,
            mouse_id_column=config.mouse_id_column,
            first_data_row=config.mouse_id_start_row,
        )
    finally:
        workbook.close()


def print_mapping(config: Config | None = None) -> None:
    """Print the configured mouse mapping as a small command-line demo."""
    metadata = inspect_template(config or Config.load())
    print(f"Sheet: {metadata.sheet_name}")
    print(f"Mapped mice: {metadata.mapped_mouse_count}")
    print(f"Used rows: {metadata.first_used_row}..{metadata.last_used_row}")
    for mouse_id, row_number in sorted(
        metadata.mouse_id_to_row.items(), key=lambda item: item[1]
    ):
        print(f"{mouse_id} -> {row_number}")


def normalize_mouse_id(value: Any) -> str | None:
    """Normalize IDs while preserving meaningful text such as leading zeroes."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        normalized = value.strip()
        return normalized or None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def _find_used_row_bounds(worksheet: Any) -> tuple[int | None, int | None]:
    """Find the first and last rows containing a non-empty value in any column."""
    used_rows = [
        cell.row
        for cell in worksheet._cells.values()
        if cell.value is not None and cell.value != ""
    ]
    return (min(used_rows), max(used_rows)) if used_rows else (None, None)


if __name__ == "__main__":
    print_mapping()
